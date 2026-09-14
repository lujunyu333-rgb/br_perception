# -*- coding: utf-8 -*-
"""
make_v1_7_docx.py — 基于 任务_v1.5.docx 生成 任务_v1.7.docx
=============================================================
保留 v1.5 全部既有章节, 仅:
  1) 把文档内"版本："段提升为 v1.7 并加版本说明
  2) 在文末追加 v1.7 修订/增补内容(规则基线修正、场地几何、新增 SLAM/导航/任务决策、感知增量、待澄清清单、排期)

用法:
    python make_v1_7_docx.py [--src 任务_v1.5.docx] [--out 任务_v1.7.docx]

依赖: python-docx (pip install python-docx)
注意: 本脚本在 2026-09-05 生成时因交互 shell(0xC0000142)无法执行, 未实测;
      若 python-docx 导入崩溃, 请先确认 pip install python-docx 与 lxml 正常。
"""

import os
import sys

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "任务_v1.5.docx")
OUT = os.path.join(ROOT, "任务_v1.7.docx")

VERSION_NOTE = (
    "**版本**：v1.7（在 v1.5 基础上优化 —— 规则引用基线改为官方英文 v1.1；"
    "场地几何按实测/确认坐标回填；新增定位与建图、导航、任务决策层三章；"
    "感知层增加事件与资源台账；修正与中文 v1.0 译本不一致的若干条款，详见文末增补）"
)

# ─────────────────────────────────────────────────────────────────────────────
# v1.7 增补内容: 以 (kind, payload) 表示
#   kind: "h1"/"h2"/"p"/"table"
#   table: (header_list, rows_list)
# ─────────────────────────────────────────────────────────────────────────────
AMEND = [
    ("h1", "v1.7 修订与增补（在 v1.5 基础上优化）"),
    ("h2", "一、版本说明"),
    ("p", "规则引用基线从中文 v1.0 译本改为官方英文 v1.1（robocon.org.cn 2026-08-28）。"
          "场地几何用确认坐标回填，缺项标注“图2待补”。"
          "本版新增定位与建图（SLAM）、导航、任务决策层三章，以承接“整条链路在这台主机规划”的目标；"
          "感知层增加“事件 + 资源台账”，为任务决策层提供数据。"),

    ("h2", "二、规则基线修正（依据英文 v1.1，P0 级）"),
    ("table", (
        ["条款", "v1.5/v1.6 依据", "v1.1 事实", "对方案的影响"],
        [
            ["重试时手持物", "全部归还初始位置", "5.1.1 归还 与 5.3.2 保留 内部冲突，须官方确认",
             "感知“重试联动”逻辑方向可能反转（现按“归还”写）"],
            ["BR 上 L1 路径", "走 3500mm 坡道", "4.3.2 仅重试时或走上 L1 楼梯时可触地",
             "坡道很可能是 TR 上转运区通道（3.4.5），BR 起步走楼梯；需澄清"],
            ["BR 用 L1 重试区", "无前提", "5.3.1 须本场已完整进入 L1 一次", "BR_RESET 状态机要记该 flag"],
            ["挪自家地球块", "未涉及", "3.5.5 无对方天空块压顶时，可挪自己已放的地球块", "决策层可“拆塔重摆”优化"],
            ["穆斯蒂卡实物", "Mikasa 排球（黄底蓝纹）", "金 futsal 球（218-165-32）",
             "视觉检测目标冲突，须确认；做双方案/现场色卡校准"],
            ["转运区交接", "简单描述", "4.4.3 BR 只可碰完全在转运区（含上方空域）内的物品；贴地时 TR 不能放转运区",
             "增加“转运区交接事件”感知"],
            ["终场不计分", "已有 release_warning", "4.6.3/8.6.1 手持/触碰 = 0 分", "决策末段“只做能稳定结束的动作”"],
        ],
    )),

    ("h2", "三、场地几何（确认坐标回填 + 图2待补）"),
    ("table", (
        ["项", "确认值", "状态"],
        [
            ["场地", "11000×11000，原点西南(0,0)，X东 Y北", "确认"],
            ["L1 平台", "[2500,2500]–[8500,8500]，中心 5500,5500，高 600", "确认"],
            ["L2 平台", "[4000,4000]–[7000,7000]，中心 5500,5500，高 900", "确认"],
            ["核心支柱", "中心 5500,5500，Ø270（L2）", "确认"],
            ["天空棋盘", "5×5，单格 240×240，全景 1200×1200；中心格空、12 块、颜色镜像", "格距确认；绝对中心图2未给"],
            ["建筑位", "500×500", "尺寸确认；行列/坐标图2未给"],
            ["转运区", "1000×1000", "尺寸确认；中心图2未给"],
            ["启动区", "700×700（TR/BR 各一）", "尺寸确认；中心图2未给"],
            ["存储区", "1000×2000", "尺寸确认；中心图2未给"],
            ["穆斯蒂卡柱", "高 500、Ø270", "尺寸确认；位置图2未给"],
        ],
    )),
    ("p", "等轴测图显示：天空棋盘在前右地面一带、穆斯蒂卡柱在后左地面、坡道在左侧、启动区在前下方——均非几何中心，坐标以图2为准。"),

    ("h2", "四、新增：定位与建图（SLAM）要点"),
    ("p", "场地固定已知 → 不做探索式建图，做定位：FAST-LIO2 里程计 + 场地先验重定位（柱/建筑位 PnP 对齐）"
          "+ BR_RESET 位姿注入。"),
    ("p", "地图表示：每层独立 2D costmap（L1、L2）+ 层间转移边（楼梯口、L1 重试下放点）+ L1 边缘禁入带（防坠落）。"
          "依赖 OccupancyGrid；不做 3D ESDF/Voxblox（斜坡对轮式 BR 是机械限制，交给决策层）。"
          "退化处理与 ground_segmenter 坡上状态、odometry_reliable 联动（沿用 v1.5）。"),

    ("h2", "五、新增：导航要点"),
    ("p", "全局：A*/Smac 在分层 costmap 上 + 层间图（不是一张 2D 平面连续规划）。"
          "局部：DWB / RPP（低速慢平台）；不引入 MPPI（避免与感知推理抢 CPU 核）。"
          "语义化：静态障碍来自场地几何（已建成塔=障碍，空地=可行），动态障碍来自融合语义目标。"
          "与 STM32 主控边界待定：若导航/决策放主机，协议需扩展（目标点/期望路径下发）。"),

    ("h2", "六、新增：任务决策层（“任务优化”核心）"),
    ("p", "目标：180s 内最大化得分。转移 5/块；塔 L1 10/20/40、L2 20/40/80；穆斯蒂卡 250（需圣所达成 ≥2 完整塔且 ≥1 在共享区）。"),
    ("p", "结构：状态估计 → 价值评估（期望得分/风险）→ 行为调度（优先级）→ 执行。"),
    ("p", "关键输入（感知需补）：① 转运区交接台账（块进入/被取走、类型/颜色/时间 → TR 供给节奏）；"
          "② 对手威胁（对方 BR 在共享区塔边停留/翻顶）；③ 资源台账（剩余地球块、天空块朝向）；"
          "④ 圣所使命状态与风险（达成前共享区塔被翻会推迟使命）；⑤ 终场 5 秒纪律（宁可不放，也不要在蜂鸣时手持/碰块）。"),
    ("p", "风险：区域/推移违规 → 强制重试耗时间；终场手持 = 0 分。"),

    ("h2", "七、感知增量（为决策层供数）"),
    ("p", "transfer_zone 台账字段（物品进入/取走/类型/颜色/时间）；Obstacle 增加“对己方塔威胁”标记 + “塔被触碰中”状态；"
          "剩余资源：存储区地球块计数（随 TR 取走变化）、天空块朝向跟踪；天空棋盘 12 块先验（已可按 240 格计算）→ 融合层跟踪；"
          "BR_RESET 增加“首次进入 L1”flag；穆斯蒂卡竞态监控（对手 TR 靠近柱时告警）。"),

    ("h2", "八、待澄清清单（P0，找裁判/组织方确认）"),
    ("p", "1) 重试时手持物：保留（5.3.2）还是归还（5.1.1）？ 2) BR 上 L1 合法路径（楼梯 or 坡道）；3500mm 坡道归属？ "
          "3) 转运区在 L1 哪条边、TR 怎么上去（3.4.5）？ 4) 天空棋盘/穆斯蒂卡柱/建筑位/转运/启动/存储的绝对坐标（图2）。 "
          "5) 穆斯蒂卡实物（金球 or 排球）。"),

    ("h2", "九、模块优先级建议（对应 v1.5 §13.3 排期修订）"),
    ("p", "Phase 0：规则澄清 + 场地坐标回填（field_geometry.yaml）。"
          "Phase 1：感知管线（已有 9/28 模块）+ 转运台账/资源跟踪。"
          "Phase 1.5：热管理接口（沿用手动回调框架）。"
          "Phase 2：定位（FAST-LIO2 + 先验）+ 多平面 costmap + 局部导航。"
          "Phase 3：天空顶色/穆斯蒂卡（实物待澄清）+ 任务决策层骨架。"
          "Phase 4：融合 + 决策层细化 + 协议扩展（若导航/决策在主机）。"),
]


def bump_version(doc):
    """把文档中‘**版本**’段替换为 v1.7。"""
    for para in doc.paragraphs:
        if "**版本**" in para.text or para.text.strip().startswith("**版本"):
            para.text = VERSION_NOTE
            return True
    return False


def _heading(doc, text, size):
    """模板无 Heading 样式, 用普通段落手动加粗放大。"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    return p


def _set_table_borders(table):
    """模板无 Table Grid 样式, 用 XML 给表格加单线边框。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def append_amend(doc, items):
    for kind, payload in items:
        if kind == "h1":
            _heading(doc, payload, 16)
        elif kind == "h2":
            _heading(doc, payload, 13)
        elif kind == "p":
            doc.add_paragraph(payload)
        elif kind == "table":
            header, rows = payload
            t = doc.add_table(rows=1, cols=len(header))
            _set_table_borders(t)
            for i, h in enumerate(header):
                t.rows[0].cells[i].text = h
            for r in rows:
                cells = t.add_row().cells
                for i, c in enumerate(r):
                    cells[i].text = c
        else:
            raise ValueError("unknown kind: %s" % kind)


def main():
    src = sys.argv[sys.argv.index("--src") + 1] if "--src" in sys.argv else SRC
    out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else OUT
    doc = Document(src)
    bumped = bump_version(doc)
    print("version bumped:", bumped)
    append_amend(doc, AMEND)
    doc.save(out)
    print("saved:", out)


if __name__ == "__main__":
    main()
